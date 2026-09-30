# Servicio de gestion de balanza de camiones, importacion/exportacion de pesadas e integracion con existencias
import io
import csv
import datetime
# Importa expresiones regulares para limpieza de etiquetas y prefijos de namespace XML
import re
# Intento protegido de importacion de la libreria openpyxl para manipular planillas Excel
try:
    # Importa el modulo openpyxl
    import openpyxl
    # Importa clases de formato y estilo para celdas Excel
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
# Captura la ausencia de openpyxl en entornos restringidos
except ImportError:
    # Asigna None para permitir la carga del servicio sin dependencias obligatorias
    openpyxl = None
# Intento protegido de importacion de la libreria xlrd para planillas Excel clasicas (.xls)
try:
    # Importa el modulo xlrd
    import xlrd
# Captura la ausencia de xlrd en entornos restringidos
except ImportError:
    # Asigna None si la libreria no esta disponible
    xlrd = None
# Importa el analizador HTML de la biblioteca estandar para procesar planillas .xls formateadas como HTML
from html.parser import HTMLParser
# Importa ElementTree de la biblioteca estandar para procesar planillas .xls en formato XML Spreadsheet
import xml.etree.ElementTree as ET
from core.database import get_db_connection
from core.timezone import get_plant_now_str
from core.audit import record_audit_event
from core.error_logger import log_error, log_info

# Analizador robusto de tablas HTML para archivos generados por balanzas industriales exportados con extension .xls
class HTMLTableParser(HTMLParser):
    # Inicializa el analizador de filas y celdas
    def __init__(self):
        # Llama al constructor de la clase base
        super().__init__()
        # Lista de filas acumuladas
        self.rows = []
        # Fila actual en proceso
        self.current_row = []
        # Contenido de la celda actual
        self.current_cell = []
        # Bandera de celda abierta
        self.in_cell = False

    # Procesa apertura de etiquetas html
    def handle_starttag(self, tag, attrs):
        # Normaliza la etiqueta a minusculas
        tag_l = tag.lower()
        # Si es celda de datos o encabezado
        if tag_l in ('td', 'th'):
            # Activa bandera de celda
            self.in_cell = True
            # Reinicia el acumulador de texto de celda
            self.current_cell = []
        # Si es inicio de fila
        elif tag_l == 'tr':
            # Reinicia la fila actual
            self.current_row = []

    # Procesa texto dentro de las etiquetas
    def handle_data(self, data):
        # Si se encuentra dentro de una celda
        if self.in_cell:
            # Acumula el fragmento de texto
            self.current_cell.append(data)

    # Procesa cierre de etiquetas html
    def handle_endtag(self, tag):
        # Normaliza etiqueta a minusculas
        tag_l = tag.lower()
        # Si se cierra una celda
        if tag_l in ('td', 'th'):
            # Desactiva bandera de celda
            self.in_cell = False
            # Concatena y limpia el texto de la celda
            cell_text = ''.join(self.current_cell).strip()
            # Agrega el texto a la fila actual
            self.current_row.append(cell_text)
        # Si se cierra una fila
        elif tag_l == 'tr':
            # Si la fila tiene al menos una celda con contenido
            if self.current_row:
                # Agrega la fila al conjunto de filas extraidas
                self.rows.append(self.current_row)

# Normaliza texto para comparacion insensible a acentos, espacios y mayusculas
def _clean_header(h):
    if not h:
        return ''
    h = str(h).strip().lower()
    h = h.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
    h = h.replace('_', ' ').replace('-', ' ').replace('.', '')
    return ' '.join(h.split())

# Determina el tipo de operacion y producto estandarizado a partir de texto libre
def _normalize_operation_and_product(raw_op, raw_prod):
    op = (raw_op or '').strip().lower()
    prod = (raw_prod or '').strip().lower()

    # Normalizacion de operacion
    if any(k in op for k in ['egreso', 'salida', 'despacho', 'venta', 'embarque']):
        operation_type = 'egreso'
    elif any(k in op for k in ['ingreso', 'entrada', 'compra', 'recepcion', 'descarga']):
        operation_type = 'ingreso'
    elif any(k in op for k in ['interno', 'trasvase', 'recirculacion']):
        operation_type = 'interno'
    else:
        # Deduccion segun el producto si la operacion no esta explicita
        if 'semilla' in prod or 'girasol' in prod:
            operation_type = 'ingreso'
        elif 'aceite' in prod or 'expeller' in prod:
            operation_type = 'egreso'
        else:
            operation_type = 'ingreso'

    # Normalizacion de producto
    if 'semilla' in prod or 'grano' in prod or 'girasol' in prod:
        product_norm = 'semilla'
    elif 'aceite' in prod:
        product_norm = 'aceite'
    elif 'expeller' in prod or 'pellet' in prod or 'harina' in prod:
        product_norm = 'expeller'
    elif 'insumo' in prod or 'repuesto' in prod or 'quimico' in prod or 'solvente' in prod or 'lena' in prod or 'gas' in prod:
        product_norm = 'insumos'
    else:
        product_norm = prod or 'insumos'

    return operation_type, product_norm

# Registra una pesada de balanza de camion en la base de datos
def record_weighing(
    ticket_number=None, weigh_date=None, operation_type='ingreso', product='semilla',
    truck_plate=None, trailer_plate=None, transport_company=None,
    driver_name=None, driver_dni=None,
    gross_weight_kg=0.0, tare_weight_kg=0.0, net_weight_kg=None,
    origin=None, destination=None, origin_name=None, destination_name=None,
    seals_numbers=None, notes=None,
    operator_name='Balanza', shift_id='TC', sync_inventory=False
):
    origin = origin or origin_name
    destination = destination or destination_name
    gross = float(gross_weight_kg or 0.0)
    tare = float(tare_weight_kg or 0.0)

    # Si no se envio neto explicito, se calcula por diferencia
    if net_weight_kg is not None and float(net_weight_kg) > 0:
        net = float(net_weight_kg)
    else:
        net = max(0.0, gross - tare)

    net_tons = round(net / 1000.0, 3)

    if not weigh_date:
        weigh_date = get_plant_now_str()

    with get_db_connection() as conn:
        cursor = conn.execute("""
            INSERT INTO truck_scale_weighings (
                ticket_number, weigh_date, operation_type, product,
                truck_plate, trailer_plate, transport_company,
                driver_name, driver_dni,
                gross_weight_kg, tare_weight_kg, net_weight_kg, net_weight_tons,
                origin, destination, seals_numbers, notes,
                operator_name, shift_id, sync_inventory
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            ticket_number, weigh_date, operation_type, product,
            (truck_plate or '').upper(), (trailer_plate or '').upper() if trailer_plate else None,
            transport_company, driver_name, driver_dni,
            gross, tare, net, net_tons,
            origin, destination, seals_numbers, notes,
            operator_name, shift_id, 1 if sync_inventory else 0
        ))
        weighing_id = cursor.lastrowid

        # Si se solicita sincronizar con existencias de inventario
        if sync_inventory and net > 0:
            from modules.inventory.service import record_inventory_movement
            try:
                # El movimiento de inventario impacta en el stock general
                inv_mov_type = 'ingreso' if operation_type == 'ingreso' else 'despacho'
                inv_prod = product if product in ['semilla', 'aceite', 'expeller'] else 'semilla'
                inv_notes = f"Balanza Ticket {ticket_number or weighing_id} - Camión {truck_plate or '-'}"
                record_inventory_movement(
                    product=inv_prod,
                    movement_type=inv_mov_type,
                    origin=origin or 'Balanza Báscula',
                    destination=destination or 'Planta BioBalcarce',
                    quantity_kg=net,
                    document_ref=f"BAL-{ticket_number or weighing_id}",
                    shift_id=shift_id,
                    operator_name=operator_name,
                    notes=inv_notes
                )
            except Exception as e:
                log_error('WEIGHBRIDGE_SERVICE', f"No se pudo sincronizar inventario para pesada {weighing_id}", e)

        conn.commit()

    return weighing_id

# Obtiene la lista filtrada de pesadas de balanza para la vista web
def get_recent_weighings(limit=200, product=None, operation_type=None, search=None, start_date=None, end_date=None):
    query = "SELECT * FROM truck_scale_weighings WHERE 1=1"
    params = []

    if product and product != 'todos':
        query += " AND product = ?"
        params.append(product)

    if operation_type and operation_type != 'todos':
        query += " AND operation_type = ?"
        params.append(operation_type)

    if search:
        term = f"%{search.strip()}%"
        query += " AND (truck_plate LIKE ? OR trailer_plate LIKE ? OR ticket_number LIKE ? OR driver_name LIKE ? OR transport_company LIKE ? OR notes LIKE ?)"
        params.extend([term, term, term, term, term, term])

    if start_date:
        query += " AND weigh_date >= ?"
        params.append(f"{start_date} 00:00:00")

    if end_date:
        query += " AND weigh_date <= ?"
        params.append(f"{end_date} 23:59:59")

    query += " ORDER BY weigh_date DESC, id DESC"
    if limit:
        query += " LIMIT ?"
        params.append(limit)

    with get_db_connection() as conn:
        rows = conn.execute(query, params).fetchall()
        return [dict(r) for r in rows]

# Obtiene metricas agregadas para las tarjetas superiores de KPI
def get_weighing_summary_stats(start_date=None, end_date=None):
    query = "SELECT operation_type, product, SUM(net_weight_kg) as total_kg, COUNT(*) as count FROM truck_scale_weighings WHERE 1=1"
    params = []
    if start_date:
        query += " AND weigh_date >= ?"
        params.append(f"{start_date} 00:00:00")
    if end_date:
        query += " AND weigh_date <= ?"
        params.append(f"{end_date} 23:59:59")
    query += " GROUP BY operation_type, product"

    with get_db_connection() as conn:
        rows = conn.execute(query, params).fetchall()

    stats = {
        'total_trucks': 0,
        'total_ingreso_tons': 0.0,
        'total_egreso_tons': 0.0,
        'semilla_ingreso_tons': 0.0,
        'expeller_egreso_tons': 0.0,
        'aceite_egreso_tons': 0.0,
        'insumos_tons': 0.0
    }

    for r in rows:
        op = r['operation_type']
        prod = r['product']
        tons = round((r['total_kg'] or 0.0) / 1000.0, 2)
        count = r['count'] or 0

        stats['total_trucks'] += count
        if op == 'ingreso':
            stats['total_ingreso_tons'] += tons
            if prod == 'semilla':
                stats['semilla_ingreso_tons'] += tons
            else:
                stats['insumos_tons'] += tons
        elif op == 'egreso':
            stats['total_egreso_tons'] += tons
            if prod == 'expeller':
                stats['expeller_egreso_tons'] += tons
            elif prod == 'aceite':
                stats['aceite_egreso_tons'] += tons
            else:
                stats['insumos_tons'] += tons

    # Aliases de compatibilidad
    stats['seed_in_tons'] = stats['semilla_ingreso_tons']
    stats['oil_out_tons'] = stats['aceite_egreso_tons']
    stats['expeller_out_tons'] = stats['expeller_egreso_tons']

    return stats

# Extrae renglones crudos soportando XLSX moderno, XLS binario, XLS HTML, XLS XML y CSV/TSV
def _extract_rows_from_file(raw_bytes, filename=''):
    # Inicializa lista de filas extraidas
    raw_rows = []
    # Normaliza el nombre de archivo a minusculas
    fn = str(filename).lower()

    # 1. Intento A: Si empieza con firma ZIP (PK\x03\x04) o tiene extension .xlsx, intenta con openpyxl
    if raw_bytes.startswith(b'PK\x03\x04') or fn.endswith('.xlsx'):
        # Si openpyxl esta disponible en el entorno
        if openpyxl is not None:
            # Captura posibles excepciones de openpyxl
            try:
                # Carga el libro Excel con solo valores de datos
                wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
                # Selecciona la hoja activa
                sheet = wb.active
                # Itera sobre las filas de la hoja
                for row in sheet.iter_rows(values_only=True):
                    # Agrega la fila como lista
                    raw_rows.append(list(row))
                # Si obtuvo datos validos, los retorna de inmediato
                if raw_rows and len(raw_rows) >= 2:
                    # Retorna las filas extraidas
                    return raw_rows
            # Si falla openpyxl (por ejemplo si el archivo no es un zip real)
            except Exception:
                # Continua con los siguientes analizadores alternativos
                pass

    # 2. Intento B: Si empieza con firma OLE2 (Compound File \xd0\xcf\x11\xe0) o tiene extension .xls binaria
    if raw_bytes.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1') or fn.endswith('.xls'):
        # Si xlrd esta disponible
        if xlrd is not None:
            # Bloque de captura de errores de formato binario
            try:
                # Abre el libro binario con xlrd
                wb = xlrd.open_workbook(file_contents=raw_bytes)
                # Selecciona la primera hoja de calculo
                sheet = wb.sheet_by_index(0)
                # Itera por el numero de filas de la hoja
                for r in range(sheet.nrows):
                    # Lista de valores de la fila
                    row_vals = []
                    # Itera sobre cada columna
                    for c in range(sheet.ncols):
                        # Obtiene el tipo de celda
                        c_type = sheet.cell_type(r, c)
                        # Obtiene el valor crudo de la celda
                        c_val = sheet.cell_value(r, c)
                        # Si la celda es de tipo fecha
                        if c_type == xlrd.XL_CELL_DATE:
                            # Intenta convertir la fecha Excel a formato estandar
                            try:
                                # Convierte a tupla de fecha
                                dt_tuple = xlrd.xldate_as_tuple(c_val, wb.datemode)
                                # Crea el objeto datetime
                                dt = datetime.datetime(*dt_tuple)
                                # Agrega fecha formateada
                                row_vals.append(dt.strftime('%Y-%m-%d %H:%M:%S'))
                            # En caso de error en la fecha
                            except Exception:
                                # Guarda representacion textual
                                row_vals.append(str(c_val))
                        # Si la celda es de tipo numerico
                        elif c_type == xlrd.XL_CELL_NUMBER:
                            # Si es un numero entero
                            if isinstance(c_val, float) and c_val.is_integer():
                                # Agrega como entero
                                row_vals.append(int(c_val))
                            # Si tiene decimales
                            else:
                                # Agrega el numero flotante
                                row_vals.append(c_val)
                        # Para otros tipos de celdas (texto, booleano, etc.)
                        else:
                            # Agrega el valor directamente
                            row_vals.append(c_val)
                    # Agrega la fila a la coleccion
                    raw_rows.append(row_vals)
                # Si obtuvo datos validos, los retorna
                if raw_rows and len(raw_rows) >= 2:
                    # Retorna las filas
                    return raw_rows
            # Si no es un archivo BIFF8 soportado
            except Exception:
                # Continua con otros analizadores
                pass

    # 3. Intento C: Decodificacion textual (HTML Table, SpreadsheetML XML, CSV o TSV)
    # Intenta decodificar con distintos encodings comunes en software de balanzas
    text_content = None
    # Lista de codificaciones a probar
    for enc in ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']:
        # Intenta decodificar
        try:
            # Decodifica los bytes
            text_content = raw_bytes.decode(enc)
            # Corta el bucle si tuvo exito
            break
        # Captura error de decodificacion
        except (UnicodeDecodeError, Exception):
            # Prueba con el siguiente encoding
            continue

    # Si se logro decodificar a texto
    if text_content:
        # Muestra en minusculas para deteccion de etiquetas
        sample = text_content[:4096].lower()

        # C1. Analisis si es una tabla HTML exportada con extension .xls
        if '<table' in sample or '<tr' in sample or '<html' in sample:
            # Intento de parseo con HTMLTableParser
            try:
                # Instancia el analizador HTML
                html_p = HTMLTableParser()
                # Procesa el contenido completo
                html_p.feed(text_content)
                # Si encontro filas validas
                if html_p.rows and len(html_p.rows) >= 2:
                    # Retorna las filas de la tabla HTML
                    return html_p.rows
            # Si ocurre error al parsear HTML
            except Exception:
                # Continua con las demas opciones
                pass

        # C2. Analisis si es un SpreadsheetML XML exportado como .xls
        if '<workbook' in sample or '<?xml' in sample:
            # Intento de parseo de XML Spreadsheet
            try:
                # Limpia prefijos de namespace de etiquetas como ss: para evitar errores de unbound prefix
                cleaned_xml = re.sub(r'<(/?)(\w+):', r'<\1', text_content)
                # Limpia prefijos de atributos con namespace como ss:Name o ss:Type
                cleaned_xml = re.sub(r'\s+(\w+):(\w+)=', r' \2=', cleaned_xml)
                # Parsea el arbol XML saneado
                root = ET.fromstring(cleaned_xml)
                # Lista contenedora de filas XML extraidas
                xml_rows = []
                # Itera buscando elementos de fila en el documento
                for elem in root.iter():
                    # Si el tag termina en Row (insensible al namespace)
                    if elem.tag.endswith('Row'):
                        # Celdas de la fila actual
                        row_cells = []
                        # Itera sobre los elementos hijos directos de la fila
                        for cell_elem in elem:
                            # Valor por defecto de la celda
                            cell_val = ''
                            # Itera sobre los elementos contenidos en la celda
                            for child in cell_elem:
                                # Si el elemento es un contenedor de datos Data
                                if child.tag.endswith('Data'):
                                    # Extrae el texto del dato
                                    cell_val = child.text or ''
                                    # Finaliza la busqueda de dato para esta celda
                                    break
                            # Si no tenia hijo Data pero tenia texto directo en la celda
                            if not cell_val and cell_elem.text and cell_elem.text.strip():
                                # Asigna el texto directo encontrado
                                cell_val = cell_elem.text.strip()
                            # Agrega el valor de texto a la fila
                            row_cells.append(cell_val)
                        # Si la fila contiene al menos una celda
                        if row_cells:
                            # Agrega la fila a la coleccion de filas
                            xml_rows.append(row_cells)
                # Si extrajo al menos fila de cabecera y una fila de datos
                if xml_rows and len(xml_rows) >= 2:
                    # Retorna las filas de la planilla XML
                    return xml_rows
            # Si falla el parseo XML
            except Exception:
                # Continua con la siguiente estrategia
                pass

        # C3. Analisis si es texto delimitado (CSV / TSV / punto y coma)
        # Cuenta ocurrencias de delimitadores en la muestra
        cnt_tab = sample.count('\t')
        # Cuenta punto y coma
        cnt_semi = sample.count(';')
        # Cuenta comas
        cnt_comma = sample.count(',')
        # Determina el delimitador mas probable
        if cnt_tab > cnt_semi and cnt_tab > cnt_comma and cnt_tab > 2:
            # Tabulador
            delim = '\t'
        elif cnt_semi > cnt_comma:
            # Punto y coma
            delim = ';'
        else:
            # Coma por defecto
            delim = ','

        # Intento de lectura CSV
        try:
            # Crea lector CSV
            csv_reader = csv.reader(io.StringIO(text_content), delimiter=delim)
            # Extrae todas las filas
            csv_rows = [list(r) for r in csv_reader if r]
            # Si obtuvo datos coherentes
            if csv_rows and len(csv_rows) >= 2:
                # Retorna las filas CSV
                return csv_rows
        # Si falla el lector CSV
        except Exception:
            # Continua
            pass

    # 4. Fallback final: Si openpyxl todavia no se intento y esta disponible
    if openpyxl is not None and not raw_bytes.startswith(b'PK\x03\x04'):
        # Intenta una ultima vez con openpyxl
        try:
            # Carga el archivo
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
            # Lee filas
            raw_rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
            # Si tiene datos
            if raw_rows and len(raw_rows) >= 2:
                # Retorna filas
                return raw_rows
        # Si falla
        except Exception:
            # Pasa
            pass

    # Si ninguno de los analizadores logro extraer los datos
    if not raw_rows or len(raw_rows) < 2:
        # Lanza error descriptivo indicando las alternativas compatibles
        raise ValueError(
            "No fue posible interpretar el archivo de balanza. El sistema admite Excel moderno (.xlsx), "
            "Excel clásico (.xls en formato binario, HTML o XML) y texto delimitado (.csv). "
            "Verifique que el archivo contenga datos o guarde la planilla en formato .xlsx o .csv antes de subirla."
        )

    # Retorna las filas
    return raw_rows

# Importador inteligente de archivos Excel (.xlsx, .xls) y CSV de balanza
def import_weighings_from_file(file_storage, filename=None, operator_name='Balanza', shift_id='TC', sync_inventory=False):
    # Si no se envio el nombre del archivo, intenta obtenerlo del objeto file_storage
    if filename is None:
        # Obtiene el nombre del archivo adjunto
        filename = getattr(file_storage, 'filename', '') or ''
    # Normaliza el nombre del archivo a minusculas
    filename = str(filename).lower()
    # Inicializa la lista de registros procesados
    records = []

    # Extrae el contenido en bytes del archivo recibido
    if isinstance(file_storage, bytes):
        # Asigna directamente si ya son bytes
        raw_bytes = file_storage
    # Si es un objeto tipo stream o archivo de Flask
    elif hasattr(file_storage, 'read'):
        # Lee los bytes del flujo
        raw_bytes = file_storage.read()
        # Si tiene soporte para rebobinar el cursor, lo reinicia
        if hasattr(file_storage, 'seek'):
            # Vuelve el puntero al inicio
            file_storage.seek(0)
    # En caso de cadenas de texto u otros tipos
    else:
        # Convierte a representacion en bytes
        raw_bytes = str(file_storage).encode('utf-8')

    # Extrae las filas del archivo utilizando el extractor multitipo universal
    raw_rows = _extract_rows_from_file(raw_bytes, filename=filename)

    # Verifica que el archivo no este vacio o contenga un solo renglon
    if not raw_rows or len(raw_rows) < 2:
        # Lanza error si no hay datos suficientes
        raise ValueError("El archivo se encuentra vacío o no posee renglones de datos.")

    # Busca la fila de encabezados analizando las primeras 10 lineas
    header_idx = -1
    col_map = {}

    for idx, row in enumerate(raw_rows[:10]):
        if not row:
            continue
        cleaned = [_clean_header(c) for c in row if c is not None]
        # Si la fila contiene palabras clave de balanza
        keywords = ['ticket', 'fecha', 'patente', 'bruto', 'tara', 'neto', 'producto', 'chofer', 'camion', 'material']
        match_count = sum(1 for kw in keywords if any(kw in c for c in cleaned))
        if match_count >= 2:
            header_idx = idx
            break

    if header_idx == -1:
        # Fallback al primer renglon si no se detectaron encabezados tipicos
        header_idx = 0

    header_row = raw_rows[header_idx]
    for c_idx, cell in enumerate(header_row):
        norm = _clean_header(cell)
        if not norm:
            continue

        # Mapeo de columnas
        if any(k in norm for k in ['ticket', 'comprobante', 'remito', 'boleta', 'nro']):
            col_map['ticket'] = c_idx
        elif any(k in norm for k in ['fecha', 'hora', 'date', 'timestamp']):
            if 'fecha' not in col_map: # Prioriza primer campo fecha
                col_map['fecha'] = c_idx
        elif any(k in norm for k in ['chasis', 'patente camion', 'dominio chasis', 'camion', 'patente']):
            if 'chasis' not in col_map:
                col_map['chasis'] = c_idx
        elif any(k in norm for k in ['acoplado', 'semi', 'trailer', 'patente acoplado', 'dominio acoplado']):
            col_map['acoplado'] = c_idx
        elif any(k in norm for k in ['chofer', 'conductor', 'transportista']):
            col_map['chofer'] = c_idx
        elif any(k in norm for k in ['dni', 'cuit', 'documento']):
            col_map['dni'] = c_idx
        elif any(k in norm for k in ['empresa', 'transporte', 'razon social']):
            col_map['transporte'] = c_idx
        elif any(k in norm for k in ['bruto', 'peso bruto', 'kg bruto', 'pesada 1']):
            col_map['bruto'] = c_idx
        elif any(k in norm for k in ['tara', 'peso tara', 'kg tara', 'pesada 2']):
            col_map['tara'] = c_idx
        elif any(k in norm for k in ['neto', 'peso neto', 'kg neto']):
            col_map['neto'] = c_idx
        elif any(k in norm for k in ['producto', 'material', 'cereal', 'mercaderia']):
            col_map['producto'] = c_idx
        elif any(k in norm for k in ['operacion', 'movimiento', 'sentido', 'tipo']):
            col_map['operacion'] = c_idx
        elif any(k in norm for k in ['origen', 'procedencia', 'proveedor']):
            col_map['origen'] = c_idx
        elif any(k in norm for k in ['destino', 'cliente', 'destinatario']):
            col_map['destino'] = c_idx
        elif any(k in norm for k in ['precinto', 'precintos', 'seals']):
            col_map['precintos'] = c_idx
        elif any(k in norm for k in ['observacion', 'observaciones', 'notas', 'obs']):
            col_map['notas'] = c_idx

    imported_count = 0
    skipped_count = 0
    errors = []

    # Procesa cada fila de datos a partir de header_idx + 1
    for r_idx, row in enumerate(raw_rows[header_idx + 1:], start=header_idx + 2):
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue

        def get_val(key, default=''):
            if key in col_map and col_map[key] < len(row):
                v = row[col_map[key]]
                if v is None:
                    return default
                if isinstance(v, datetime.datetime):
                    return v.strftime('%Y-%m-%d %H:%M:%S')
                if isinstance(v, datetime.date):
                    return v.strftime('%Y-%m-%d')
                return str(v).strip()
            return default

        def get_float(key, default=0.0):
            val_str = get_val(key)
            if not val_str:
                return default
            try:
                # Normaliza separadores de miles y decimales
                val_str = val_str.replace('.', '').replace(',', '.') if ',' in val_str and '.' in val_str else val_str.replace(',', '.')
                return float(val_str)
            except ValueError:
                return default

        ticket = get_val('ticket')
        fecha = get_val('fecha') or get_plant_now_str()
        chasis = get_val('chasis')
        acoplado = get_val('acoplado')
        chofer = get_val('chofer')
        dni = get_val('dni')
        transporte = get_val('transporte')
        raw_prod = get_val('producto')
        raw_op = get_val('operacion')
        bruto = get_float('bruto')
        tara = get_float('tara')
        neto = get_float('neto')
        origen = get_val('origen')
        destino = get_val('destino')
        precintos = get_val('precintos')
        notas = get_val('notas')

        # Si no hay chasis ni ticket ni peso, saltea la linea
        if not chasis and not ticket and bruto == 0.0 and neto == 0.0:
            skipped_count += 1
            continue

        # Si el neto es 0 pero hay bruto y tara, calcula
        if neto <= 0.0 and bruto > 0.0 and tara > 0.0:
            neto = max(0.0, bruto - tara)

        # Si el valor de los pesos vino en toneladas (ej. 28.5 en vez de 28500 kg)
        if 0 < neto < 150: # Evidente carga en toneladas
            neto = neto * 1000.0
            bruto = bruto * 1000.0 if bruto < 150 else bruto
            tara = tara * 1000.0 if tara < 150 else tara

        op_type, prod_norm = _normalize_operation_and_product(raw_op, raw_prod)

        try:
            record_weighing(
                ticket_number=ticket,
                weigh_date=fecha,
                operation_type=op_type,
                product=prod_norm,
                truck_plate=chasis,
                trailer_plate=acoplado,
                transport_company=transporte,
                driver_name=chofer,
                driver_dni=dni,
                gross_weight_kg=bruto,
                tare_weight_kg=tara,
                net_weight_kg=neto,
                origin=origen,
                destination=destino,
                seals_numbers=precintos,
                notes=notas,
                operator_name=operator_name,
                shift_id=shift_id,
                sync_inventory=sync_inventory
            )
            imported_count += 1
        except Exception as e:
            errors.append(f"Fila {r_idx}: {str(e)}")
            skipped_count += 1

    record_audit_event('BALANZA', 'IMPORTACION_EXCEL', f"Importación de balanza: {imported_count} registros incorporados, {skipped_count} omitidos.")
    return {
        'success': True,
        'imported_count': imported_count,
        'skipped_count': skipped_count,
        'errors': errors
    }

# Exporta los registros a un archivo Excel (.xlsx) estructurado
def export_weighings_to_excel(weighings=None, as_stream=False):
    # Verifica la disponibilidad de la libreria openpyxl
    if openpyxl is None:
        # Lanza excepcion descriptiva si no se encuentra instalada
        raise RuntimeError("La exportación a formato Excel requiere la librería openpyxl instalada.")
    if weighings is None:
        weighings = get_recent_weighings(limit=5000)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pesadas de Balanza"

    # Estilos corporativos de BioBalcarce
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Slate 900
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9)
    bold_font = Font(name="Arial", size=9, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        "ID", "Fecha / Hora", "N° Ticket", "Operación", "Producto",
        "Chasis", "Acoplado", "Transporte", "Chofer", "DNI",
        "Bruto (kg)", "Tara (kg)", "Neto (kg)", "Neto (Tn)",
        "Origen", "Destino", "Precintos", "Observaciones", "Operador"
    ]

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    ws.row_dimensions[1].height = 26

    for r_idx, w in enumerate(weighings, start=2):
        row_data = [
            w.get('id'),
            w.get('weigh_date'),
            w.get('ticket_number') or '-',
            (w.get('operation_type') or '').upper(),
            (w.get('product') or '').capitalize(),
            w.get('truck_plate') or '-',
            w.get('trailer_plate') or '-',
            w.get('transport_company') or '-',
            w.get('driver_name') or '-',
            w.get('driver_dni') or '-',
            float(w.get('gross_weight_kg') or 0.0),
            float(w.get('tare_weight_kg') or 0.0),
            float(w.get('net_weight_kg') or 0.0),
            float(w.get('net_weight_tons') or 0.0),
            w.get('origin') or '-',
            w.get('destination') or '-',
            w.get('seals_numbers') or '-',
            w.get('notes') or '-',
            w.get('operator_name') or '-'
        ]
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20

        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=r_idx, column=col_idx)
            c.font = data_font
            c.border = thin_border
            if col_idx in [11, 12, 13]: # Pesos kg
                c.number_format = '#,##0'
                c.alignment = right_align
            elif col_idx == 14: # Tn
                c.number_format = '#,##0.00'
                c.font = bold_font
                c.alignment = right_align
            elif col_idx in [1, 2, 4, 6, 7, 10]:
                c.alignment = center_align
            else:
                c.alignment = left_align

    # Autoajuste de anchos de columna
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    if as_stream:
        return output
    return output.getvalue()

# Exporta los registros a formato CSV
def export_weighings_to_csv(weighings=None):
    if weighings is None:
        weighings = get_recent_weighings(limit=5000)
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')

    headers = [
        "ID", "Fecha_Hora", "Ticket", "Operacion", "Producto",
        "Chasis", "Acoplado", "Transporte", "Chofer", "DNI",
        "Bruto_kg", "Tara_kg", "Neto_kg", "Neto_Tn",
        "Origen", "Destino", "Precintos", "Observaciones", "Operador"
    ]
    writer.writerow(headers)

    for w in weighings:
        writer.writerow([
            w.get('id'),
            w.get('weigh_date'),
            w.get('ticket_number') or '',
            w.get('operation_type') or '',
            w.get('product') or '',
            w.get('truck_plate') or '',
            w.get('trailer_plate') or '',
            w.get('transport_company') or '',
            w.get('driver_name') or '',
            w.get('driver_dni') or '',
            w.get('gross_weight_kg') or 0.0,
            w.get('tare_weight_kg') or 0.0,
            w.get('net_weight_kg') or 0.0,
            w.get('net_weight_tons') or 0.0,
            w.get('origin') or '',
            w.get('destination') or '',
            w.get('seals_numbers') or '',
            w.get('notes') or '',
            w.get('operator_name') or ''
        ])

    output.seek(0)
    return output.getvalue()

# Elimina una pesada (solo admin)
def delete_weighing(weighing_id, admin_user='admin_sistema'):
    with get_db_connection() as conn:
        row = conn.execute("SELECT * FROM truck_scale_weighings WHERE id = ?;", (weighing_id,)).fetchone()
        if not row:
            raise ValueError(f"Pesada con ID {weighing_id} no encontrada.")
        conn.execute("DELETE FROM truck_scale_weighings WHERE id = ?;", (weighing_id,))
        conn.commit()
    record_audit_event('BALANZA', 'ELIMINACION_PESADA', f"Pesada ID {weighing_id} (Ticket: {row['ticket_number']}, Camión: {row['truck_plate']}) eliminada por {admin_user}.")
    return True
